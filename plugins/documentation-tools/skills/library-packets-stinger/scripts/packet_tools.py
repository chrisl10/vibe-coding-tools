#!/usr/bin/env python3
"""Local, evidence-bound library packets. Never runs source files or downloads."""
from __future__ import annotations
import argparse
import hashlib
import importlib
import ipaddress
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import urlparse
import zipfile
from xml.etree import ElementTree

SCHEMA_VERSION = 1
DEFAULT_PREPARED_BY = "Legion Code Inc."
PREPARED_BY_ENV = "LIBRARY_PACKETS_PREPARED_BY"
AUTHOR_ENV = "LIBRARY_PACKETS_AUTHOR"
SOFFICE_ENV = "SOFFICE_PATH"
SOFFICE_NAMES = {"soffice", "soffice.bin", "soffice.exe", "libreoffice"}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ZIP_BYTES = 128 * 1024 * 1024
KINDS = {"fact", "intent", "inference", "proposal"}
STATUSES = {"unknown": 0, "proposal": 1, "intent": 1, "documented": 2, "verified": 3}
PAGE_CHECKS = ("whitespace", "cutoffs", "fonts", "tables", "figures",
               "footer", "logo", "citations", "word_layout")
CITE = re.compile(r"\[([A-Za-z][A-Za-z0-9_-]*)\](?!\()")
SAFE_ID = re.compile(r"^[A-Za-z][A-Za-z0-9_-]{0,63}$")
PRIVATE_TEXT = re.compile(
    r"(?i)(?:/Users/|/home/|[A-Z]:\\|file://|localhost|127\.0\.0\.1|"
    r"\b(?:library|server|packages|src)/|\.env\b|"
    r"\b(?:api key|bearer token|SQL|JSONB|MCP|Pydantic|Drizzle|"
    r"internal endpoint|database schema|stack trace)\b)")


class PacketError(Exception):
    pass


class Blocked(PacketError):
    pass


def resolve_branding(prepared_by=None, author=None):
    """Footer preparer and DOCX author. Precedence: argument, environment, default."""
    prepared = (prepared_by or os.environ.get(PREPARED_BY_ENV) or DEFAULT_PREPARED_BY).strip()
    docx_author = (author or os.environ.get(AUTHOR_ENV) or prepared).strip()
    if not prepared or not docx_author:
        raise PacketError("Prepared-by and author text cannot be empty")
    if "\n" in prepared or "|" in prepared:
        raise PacketError("Prepared-by text must be one line without a pipe character")
    return prepared, docx_author


def find_soffice(explicit=None):
    """Locate an installed LibreOffice. Argument, SOFFICE_PATH, PATH, then common Windows installs."""
    for candidate in (explicit, os.environ.get(SOFFICE_ENV)):
        if candidate:
            return candidate
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for variable in ("ProgramFiles", "ProgramFiles(x86)", "ProgramW6432"):
        base = os.environ.get(variable)
        if base:
            path = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if path.is_file():
                return str(path)
    return None


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def json_read(path):
    p = Path(path)
    if p.stat().st_size > MAX_FILE_BYTES:
        raise PacketError("JSON input exceeds size limit")
    return json.loads(p.read_text(encoding="utf-8"))


def json_write(path, data):
    if not all(hasattr(os, flag) for flag in ("O_DIRECTORY", "O_NOFOLLOW")):
        raise Blocked("Private JSON output requires POSIX directory and no-follow support")
    p = Path(path)
    for parent in p.absolute().parents:
        if parent.is_symlink():
            raise PacketError("Symlink JSON output ancestor refused")
    p.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    parent_fd = os.open(p.parent, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        fd = os.open(p.name, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                     0o600, dir_fd=parent_fd)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
            f.write("\n")
    finally:
        os.close(parent_fd)


def root_path(base, value):
    p = Path(value)
    p = p if p.is_absolute() else Path(base) / p
    if p.is_symlink() or not p.is_dir():
        raise PacketError(f"Root must be an existing non-symlink directory: {value}")
    return p.resolve()


def safe_path(root, value, must_exist=True):
    """Input is relative to its explicit root. Reject traversal and all symlinks."""
    q = Path(value)
    if q.is_absolute() or ".." in q.parts or not q.parts:
        raise PacketError(f"Unsafe relative path: {value}")
    root = Path(root).resolve()
    p = root / q
    for i in range(1, len(q.parts) + 1):
        if (root / Path(*q.parts[:i])).is_symlink():
            raise PacketError(f"Symlink input refused: {value}")
    if not p.resolve().is_relative_to(root):
        raise PacketError(f"Path escapes root: {value}")
    if must_exist and (not p.is_file() or p.stat().st_size > MAX_FILE_BYTES):
        raise PacketError(f"Missing, non-file or oversized input: {value}")
    return p


def optional(name):
    try:
        return importlib.import_module(name)
    except ImportError as e:
        raise Blocked(f"Required installed package unavailable: {name}") from e


def extract(path):
    p = Path(path)
    if p.stat().st_size > MAX_FILE_BYTES:
        raise PacketError("Source exceeds file limit")
    suffix = p.suffix.lower()
    if suffix in {".md", ".txt", ".mdx", ".rst", ".adoc", ".csv",
                  ".json", ".yaml", ".yml"}:
        return {"text": p.read_text(encoding="utf-8"), "pages": None}
    if suffix == ".docx":
        with zipfile.ZipFile(p) as z:
            entries = z.infolist()
            if len(entries) > 5000 or sum(e.file_size for e in entries) > MAX_ZIP_BYTES:
                raise PacketError("DOCX expansion exceeds safety limit")
            raw = z.read("word/document.xml")
            if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
                raise PacketError("DOCX XML declarations refused")
            doc = ElementTree.fromstring(raw)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = ["".join(t.text or "" for t in n.findall(".//w:t", ns))
                          for n in doc.findall(".//w:p", ns)]
            return {"text": "\n".join(paragraphs), "pages": None}
    if suffix == ".pdf":
        reader = optional("pypdf").PdfReader(p, strict=False)
        if reader.is_encrypted:
            raise PacketError("Encrypted PDF requires a separately authorized extraction")
        if len(reader.pages) > 2000:
            raise PacketError("PDF page count exceeds safety limit")
        pages = [page.extract_text() or "" for page in reader.pages]
        if not any(s.strip() for s in pages):
            raise PacketError("No extractable PDF text; OCR/manual source review required")
        return {"text": "\n".join(pages), "pages": pages}
    raise PacketError(f"Unsupported source type: {suffix or 'none'}")


def inventory(library, excludes=()):
    root = root_path(Path.cwd(), library)
    excluded = [safe_path(root, x, False).resolve() for x in excludes]
    rows = []
    # os.walk does not traverse symlink directories. Record them as errors.
    for current, dirs, files in os.walk(root, followlinks=False):
        dirs.sort()
        files.sort()
        for name in list(dirs):
            p = Path(current) / name
            if any(p.resolve().is_relative_to(e) for e in excluded):
                dirs.remove(name)
            elif p.is_symlink():
                rows.append({"path": str(p.relative_to(root)), "type": "symlink",
                             "sha256": None, "extraction_error": "Symlink directory refused"})
                dirs.remove(name)
        for name in files:
            p = Path(current) / name
            if any(p.resolve().is_relative_to(e) for e in excluded):
                continue
            row = {"path": str(p.relative_to(root)), "type": p.suffix.lower().lstrip("."),
                   "sha256": None, "characters": 0, "extraction_error": None}
            try:
                safe_path(root, row["path"])
                row["sha256"] = digest(p)
                row["characters"] = len(extract(p)["text"])
            except (PacketError, OSError, UnicodeError, ValueError, zipfile.BadZipFile) as e:
                row["extraction_error"] = str(e)
            rows.append(row)
    return {"schema_version": 1, "library_root": str(root), "sources": rows,
            "files": len(rows), "errors": sum(bool(r["extraction_error"]) for r in rows)}


def normalized(s):
    return " ".join(s.split())


def located_text(extracted, loc):
    typ, start, end = loc.get("type"), loc.get("start"), loc.get("end")
    if type(start) is not int or type(end) is not int:
        raise PacketError("Evidence location requires integer start/end")
    text = extracted["text"]
    if typ == "characters":
        if not 0 <= start < end <= len(text):
            raise PacketError("Character location out of bounds")
        return text[start:end]
    if typ in {"lines", "pages"}:
        seq = text.splitlines() if typ == "lines" else extracted["pages"]
        if seq is None or not 1 <= start <= end <= len(seq):
            raise PacketError(f"{typ} location unavailable or out of bounds")
        return "\n".join(seq[start - 1:end])
    raise PacketError("Location type must be lines, characters or pages")


def private_leaks(text):
    hits = [m.group() for m in PRIVATE_TEXT.finditer(text)]
    for url in re.findall(r"https?://[^\s)<]+", text):
        parsed = urlparse(url)
        h = parsed.hostname or ""
        if parsed.username or parsed.password or re.search(
                r"(?i)(?:^|&)(?:api_?key|token|secret|signature|auth)=", parsed.query):
            hits.append("credential-bearing URL")
        try:
            if not ipaddress.ip_address(h).is_global:
                hits.append(h)
        except ValueError:
            if (h.endswith((".local", ".internal", ".localhost")) or "." not in h
                    or re.match(r"^(private|internal|staging|intranet)\.", h)):
                hits.append(h)
    return hits


def public_forbidden(text, packet):
    """Apply project exclusions to every emitted public field, not only prose."""
    if private_leaks(text):
        return True
    plain = normalized(text).casefold()
    tokenized = normalized(re.sub(r"[\W_]+", " ", text)).casefold()
    for exclusion in packet.get("public_exclusions", []):
        if not exclusion:
            continue
        x = normalized(exclusion).casefold()
        tx = normalized(re.sub(r"[\W_]+", " ", exclusion)).casefold()
        if x in plain or (tx and tx in tokenized):
            return True
    return False


def grade(text):
    """Flesch-Kincaid-like estimate, never a comprehension certificate."""
    clean = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    clean = CITE.sub("", clean)
    words = re.findall(r"[A-Za-z]+(?:'[A-Za-z]+)?", clean)
    if not words:
        return 0.0
    sentences = max(1, len(re.findall(r"[.!?](?:\s|$)", clean)))
    syllables = 0
    for w in words:
        w = w.lower()
        n = len(re.findall(r"[aeiouy]+", w))
        if w.endswith("e") and not w.endswith(("le", "ye")):
            n -= 1
        syllables += max(1, n)
    return round(.39 * len(words) / sentences + 11.8 * syllables / len(words) - 15.59, 1)


def parse_markdown(text, asset_root, engineering=False):
    """Explicit supported grammar; unsupported structures fail instead of vanishing."""
    out, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s == "<!-- page -->":
            out.append(("page", ""))
            i += 1
            continue
        if s == "---":
            out.append(("rule", ""))
            i += 1
            continue
        if s.startswith("<"):
            raise PacketError(f"Unsupported HTML/comment at line {i + 1}")
        if s.startswith("~~~") or s.startswith(chr(96) * 3):
            if not engineering:
                raise PacketError("Code blocks require engineering audience")
            fence = s[:3]
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith(fence):
                code.append(lines[i])
                i += 1
            if i == len(lines):
                raise PacketError("Unclosed code fence")
            out.append(("code", "\n".join(code)))
            i += 1
            continue
        m = re.fullmatch(r"!\[(.*)\]\(([^)]+)\)", s)
        if m:
            p = safe_path(asset_root, m[2])
            if p.suffix.lower() != ".png" or p.stat().st_size == 0:
                raise PacketError("Diagram assets must be nonempty local PNG files")
            out.append(("image", (m[1], p)))
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.+)$", s)
        if m:
            out.append(("heading", (len(m[1]), m[2])))
            i += 1
            continue
        if s.startswith("#"):
            raise PacketError("Only heading levels 1-4 are supported")
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                line = lines[i].strip()
                if not line.endswith("|") or r"\|" in line:
                    raise PacketError("Table needs outer pipes; escaped pipes unsupported")
                cells = [x.strip() for x in line[1:-1].split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if len(rows) < 2 or len({len(row) for row in rows}) != 1:
                raise PacketError("Table needs a header and equal-width data rows")
            out.append(("table", rows))
            continue
        m = re.match(r"^([-*+]|\d+[.)])\s+(.+)$", s)
        if m:
            out.append(("bullet" if m[1] in "-*+" else "number", m[2]))
            i += 1
            continue
        if s.startswith(">"):
            out.append(("quote", s.lstrip("> ")))
            i += 1
            continue
        if re.search(r"<[^>]+>|!\[|^\s{4}", lines[i]):
            raise PacketError(f"Unsupported Markdown structure at line {i + 1}")
        para = [s]
        i += 1
        while i < len(lines) and lines[i].strip():
            nxt = lines[i].strip()
            if re.match(r"^(#|[|>]|[-*+]\s|\d+[.)]\s|!\[|<|~~~)", nxt):
                break
            if nxt.startswith(chr(96) * 3):
                break
            if re.search(r"<[^>]+>|!\[|^\s{4}", lines[i]):
                raise PacketError(f"Unsupported Markdown structure at line {i + 1}")
            para.append(nxt)
            i += 1
        out.append(("paragraph", " ".join(para)))
    return out


def block_text(block):
    kind, data = block
    if kind == "table":
        return [" ".join(row) for row in data[1:]]
    if kind == "image":
        return [data[0]]
    if kind in {"paragraph", "bullet", "number", "quote", "code"}:
        return [data]
    return []


def validate(packet_path):
    manifest_path = Path(packet_path).resolve()
    packet = json_read(manifest_path)
    errors, warnings = [], []
    if packet.get("schema_version") != SCHEMA_VERSION:
        raise PacketError("Packet schema_version must be 1")
    grade_target = packet.get("grade_target", 8)
    if type(grade_target) not in (int, float) or not 0 < grade_target <= 8:
        raise PacketError("grade_target must be a finite number greater than 0 and at most 8")
    library = root_path(manifest_path.parent, packet["library_root"])
    manuscripts = root_path(manifest_path.parent, packet.get("manuscript_root", "."))
    assets = root_path(manifest_path.parent, packet["asset_root"])
    logo = safe_path(assets, packet.get("logo", "legion-footer.png"))
    if logo.suffix.lower() != ".png" or logo.stat().st_size == 0:
        errors.append("Real nonempty PNG footer logo is required")
    sources, extracted = {}, {}
    exclusions = [safe_path(library, x, False).resolve()
                  for x in packet.get("source_exclusions", [])]
    for src in packet.get("sources", []):
        sid = src.get("id", "")
        if not SAFE_ID.fullmatch(sid) or sid in sources:
            errors.append(f"Invalid/duplicate source ID: {sid}")
            continue
        sources[sid] = src
        try:
            f = safe_path(library, src["path"])
            if any(f.resolve().is_relative_to(e) for e in exclusions):
                raise PacketError("source is inside an excluded directory")
            if digest(f) != src.get("sha256"):
                raise PacketError("source SHA-256 mismatch")
            if src.get("kind") not in {"fact", "intent"}:
                raise PacketError("source kind must be fact or intent")
            if src.get("evidence_status") not in STATUSES:
                raise PacketError("unknown source evidence_status")
            extracted[sid] = extract(f)
        except (PacketError, OSError, ValueError, KeyError, UnicodeError, zipfile.BadZipFile) as e:
            errors.append(f"Source {sid}: {e}")
    claims = {}
    for claim in packet.get("claims", []):
        cid = claim.get("id", "")
        if not SAFE_ID.fullmatch(cid) or cid in claims:
            errors.append(f"Invalid/duplicate claim ID: {cid}")
            continue
        claims[cid] = claim
        if claim.get("kind") not in KINDS or claim.get("status") not in STATUSES:
            errors.append(f"Claim {cid}: invalid kind/status")
        review = claim.get("semantic_review", {})
        if not all(review.get(k) for k in ("reviewed_by", "at", "notes")):
            errors.append(f"Claim {cid}: explicit semantic review record required")
        evs = claim.get("evidence", [])
        if not evs:
            errors.append(f"Claim {cid}: no evidence")
        for ev in evs:
            sid = ev.get("source_id")
            if sid not in extracted:
                errors.append(f"Claim {cid}: unknown/unusable source {sid}")
                continue
            try:
                quote = ev.get("quote", "")
                if not quote.strip() or normalized(quote) not in normalized(
                        located_text(extracted[sid], ev.get("location", {}))):
                    raise PacketError("quote not found at declared location")
                src = sources[sid]
                if claim.get("kind") == "fact" and src["kind"] == "intent":
                    raise PacketError("intent source cannot be strengthened into a fact")
                if STATUSES.get(claim.get("status"), 99) > STATUSES[src["evidence_status"]]:
                    raise PacketError("claim status strengthens source status")
            except PacketError as e:
                errors.append(f"Claim {cid}: {e}")
    docs, ids = [], set()
    for doc in packet.get("documents", []):
        did = doc.get("id", "")
        if not SAFE_ID.fullmatch(did) or did in ids:
            errors.append(f"Invalid/duplicate document ID: {did}")
            continue
        ids.add(did)
        try:
            audience = doc["audience"]
            if audience not in {"general", "engineering", "public"}:
                raise PacketError("unknown audience")
            md = safe_path(manuscripts, doc["markdown"]).read_text(encoding="utf-8")
            if "\u2014" in md:
                raise PacketError("em dash found")
            blocks = parse_markdown(md, assets, audience == "engineering")
            used = set(CITE.findall(md))
            unknown = used - claims.keys()
            if unknown:
                raise PacketError("unknown citation IDs: " + ", ".join(sorted(unknown)))
            for b in blocks:
                for paragraph in block_text(b):
                    if paragraph.strip() and not CITE.search(paragraph):
                        errors.append(f"Document {did}: uncited {b[0]}: {paragraph[:80]}")
                    cids = CITE.findall(paragraph)
                    if any(claims[c].get("kind") in {"intent", "proposal"} for c in cids):
                        if re.search(r"(?i)\b(shipped|live now|available now|production.ready|"
                                     r"guaranteed|always works)\b", paragraph):
                            errors.append(f"Document {did}: apparent status strengthening")
            prose = "\n".join(t for b in blocks if b[0] != "code" for t in block_text(b))
            g = grade(prose)
            if audience != "engineering" and g > grade_target:
                errors.append(f"Document {did}: estimated grade {g} exceeds target")
            if audience == "public":
                approval = doc.get("public_approval", {})
                if approval.get("approved") is not True or not all(
                        approval.get(k) for k in ("by", "at")):
                    raise PacketError("explicit public approval record required")
                public_text = doc["title"] + "\n" + md
                # Ordinary links in a public manuscript must be approved HTTPS
                # source links. Local diagram links are copied assets, not references.
                allowed_urls = {s.get("public", {}).get("url") for s in sources.values()
                                if s.get("public", {}).get("approved") is True}
                without_images = re.sub(r"!\[.*\]\([^)]+\)", "", md)
                for link in re.findall(r"\[[^\]]+\]\(([^)]+)\)", without_images):
                    if link not in allowed_urls:
                        raise PacketError("public link is not an approved source URL")
                if public_forbidden(public_text, packet):
                    raise PacketError("public manuscript contains private/code details")
                if any(src["path"].lower() in public_text.lower()
                       for src in sources.values()):
                    raise PacketError("Public manuscript exposes a private source filename/path")
                approved_assets = {a.get("path"): a for a in packet.get("public_assets", [])}
                for kind, asset in blocks:
                    if kind == "image":
                        rel = str(asset[1].relative_to(assets))
                        approval = approved_assets.get(rel, {})
                        if (approval.get("approved") is not True or
                                approval.get("sha256") != digest(asset[1]) or
                                not all(approval.get(k) for k in ("by", "at"))):
                            raise PacketError("Public diagram requires hash-bound visual/content approval")
                        if public_forbidden(asset[1].name, packet):
                            raise PacketError("Public diagram basename leaks excluded information")
                for cid in used:
                    if public_forbidden(cid + " " + claims[cid].get("statement", ""), packet):
                        raise PacketError(f"public claim {cid} leaks private detail")
                    for ev in claims[cid]["evidence"]:
                        pub = sources[ev["source_id"]].get("public", {})
                        if pub.get("approved") is not True or not pub.get("alias"):
                            raise PacketError("approved public source alias required")
                        if public_forbidden(pub["alias"] + " " + (pub.get("url") or ""), packet):
                            raise PacketError("public reference leaks private detail")
                        if pub.get("url") and urlparse(pub["url"]).scheme != "https":
                            raise PacketError("public source URL must be HTTPS")
            docs.append({"spec": doc, "markdown": md, "blocks": blocks,
                         "claims": sorted(used), "estimated_grade": g})
        except (PacketError, OSError, UnicodeError, KeyError) as e:
            errors.append(f"Document {did}: {e}")
    if not docs or not claims or not sources:
        errors.append("Insufficient sources, claims or documents")
    warnings.append("Quote/hash/status checks do not establish semantic truth. "
                    "Semantic review records are supplied attestations, not authenticated identities.")
    return {"valid": not errors, "errors": errors, "warnings": warnings,
            "packet": packet, "documents": docs, "library": library,
            "assets": assets, "logo": logo, "manifest_path": manifest_path}


def format_locator(location):
    """Readable locator; retain the original structured location in JSON."""
    typ = location["type"]
    if typ not in {"lines", "pages", "characters"}:
        raise PacketError("Unsupported reference locator")
    start, end = location["start"], location["end"]
    if type(start) is not int or type(end) is not int:
        raise PacketError("Reference locator needs integer start/end")
    span = str(start) if start == end else f"{start}-{end}"
    return f"{typ} {span}"


def citation_projection(result, doc):
    sources = {s["id"]: s for s in result["packet"]["sources"]}
    claims = {c["id"]: c for c in result["packet"]["claims"]}
    public = doc["spec"]["audience"] == "public"
    rows, refs = [], ["# References", ""]
    for cid in doc["claims"]:
        claim = claims[cid]
        entries = []
        for ev in claim["evidence"]:
            src = sources[ev["source_id"]]
            if public:
                pub = src["public"]
                entry = {"alias": pub["alias"], "url": pub.get("url")}
                refs.append(f"- [{cid}] {pub['alias']}" +
                            (f" ({pub['url']})" if pub.get("url") else ""))
            else:
                entry = {"source_id": src["id"], "title": src["title"],
                         "path": src["path"], "sha256": src["sha256"], **ev}
                refs.append(f"- [{cid}] {src['title']}; {src['path']}; "
                            f"{format_locator(ev['location'])}")
            entries.append(entry)
        rows.append({"claim_id": cid, "kind": claim["kind"],
                     "status": claim["status"], "references": entries})
    md, projection = "\n".join(refs) + "\n", {"schema_version": 1, "claims": rows}
    if public and public_forbidden(md + json.dumps(projection, ensure_ascii=False),
                                   result["packet"]):
        raise PacketError("Public citation projection leaks excluded information")
    return md, projection


def add_inline(paragraph, text):
    # Preserve visible text and basic emphasis; links stay visible, never fetched.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|" + chr(96) + r"[^" +
                         chr(96) + r"]+" + chr(96) + r")")
    for part in pattern.split(text):
        r = paragraph.add_run(part.strip("*" + chr(96)) if pattern.fullmatch(part) else part)
        if part.startswith("**"):
            r.bold = True
        elif part.startswith("*"):
            r.italic = True
        elif part.startswith(chr(96)):
            r.font.name = "Courier New"


def write_docx(doc, markdown, logo, target, assets, prepared_by=DEFAULT_PREPARED_BY, author=None):
    dx = optional("docx")
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = dx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.8), Inches(.9)
    sec.left_margin = sec.right_margin = Inches(.8)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = d.styles[style_name]
        style.font.name = "Arial"
        fonts = style.element.rPr.rFonts
        for key in list(fonts.attrib):
            if key.lower().endswith("theme"):
                del fonts.attrib[key]
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        d.styles[name].font.name = "Arial"
        d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        color = d.styles[name].element.rPr.color
        for key in list(color.attrib):
            if key != qn("w:val"):
                del color.attrib[key]
        for border in list(d.styles[name].element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        d.styles[name].paragraph_format.keep_with_next = True
    sec.footer.is_linked_to_previous = False
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run().add_picture(str(logo), width=Inches(1.5))
    footer.add_run(f"\nPrepared by {prepared_by} | Page ").font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    for kind, data in parse_markdown(markdown, assets, doc["spec"]["audience"] == "engineering"):
        if kind == "heading":
            para = d.add_paragraph(style="Title") if data[0] == 1 else d.add_heading(level=data[0])
            add_inline(para, data[1])
        elif kind == "page":
            d.add_page_break()
        elif kind == "rule":
            d.add_paragraph()
        elif kind == "table":
            table = d.add_table(rows=1, cols=len(data[0]))
            table.style = "Table Grid"
            table.autofit = False
            for col in table.columns:
                col.width = Inches(6.9 / len(data[0]))
            for j, txt in enumerate(data[0]):
                cell = table.rows[0].cells[j]
                add_inline(cell.paragraphs[0], txt)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EFEFEF")
                cell._tc.get_or_add_tcPr().append(shade)
            trpr = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            trpr.append(repeat)
            for row in data[1:]:
                cells = table.add_row().cells
                for j, txt in enumerate(row):
                    add_inline(cells[j].paragraphs[0], txt)
            for row in table.rows:
                for cell in row.cells:
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        item = OxmlElement("w:" + side)
                        item.set(qn("w:w"), "80")
                        item.set(qn("w:type"), "dxa")
                        margins.append(item)
                    cell._tc.get_or_add_tcPr().append(margins)
        elif kind == "image":
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(data[1]), width=Inches(6.5))
            add_inline(d.add_paragraph(style="Caption"), data[0])
        else:
            style = {"bullet": "List Bullet", "number": "List Number", "quote": "Quote"}.get(kind)
            para = d.add_paragraph(style=style)
            if kind == "code":
                run = para.add_run(data)
                run.font.name, run.font.size = "Courier New", Pt(9)
            else:
                add_inline(para, data)
    d.core_properties.title = doc["spec"]["title"]
    d.core_properties.author = author or prepared_by
    d.save(target)


def inspect_pdf(pdf, docx, folder, prepared_by=DEFAULT_PREPARED_BY):
    pi = optional("pypdfium2")
    pdfdoc = pi.PdfDocument(str(pdf))
    pages = []
    for i in range(len(pdfdoc)):
        page = pdfdoc[i]
        w, h = page.get_size()
        tp = page.get_textpage()
        txt = tp.get_text_range()
        outside = 0
        for n in range(tp.count_chars()):
            if not tp.get_text_range(n, 1).strip():
                continue
            x0, y0, x1, y1 = tp.get_charbox(n)
            outside += int(x0 < -.5 or y0 < -.5 or x1 > w + .5 or y1 > h + .5)
        footer = tp.get_text_bounded(0, 0, w, 65)
        images = [obj.get_bounds() for obj in page.get_objects(max_depth=5)
                  if obj.type == pi.raw.FPDF_PAGEOBJ_IMAGE]
        logos = [b for b in images if b[1] < 80]
        outside_images = sum(x0 < -.5 or y0 < -.5 or x1 > w + .5 or y1 > h + .5
                             for x0, y0, x1, y1 in images)
        png = folder / f"page-{i + 1:03}.png"
        page.render(scale=1.5).to_pil().save(png)
        remaining_footer = re.sub(r"Prepared by\s+" + re.escape(normalized(prepared_by)) + r"\s*\|\s*Page\s*\d+", "",
                                  normalized(footer)).strip()
        pages.append({"page": i + 1, "png": png.name, "png_sha256": digest(png),
                      "out_of_page_glyphs": outside,
                      "out_of_page_images": outside_images,
                      "footer_text_present": f"Prepared by {prepared_by}" in footer,
                      "footer_logo_detected": bool(logos),
                      "footer_region_extra_text": remaining_footer,
                      "missing_glyphs": "\ufffd" in txt,
                      "heuristic_only": True})
        tp.close()
        page.close()
    pdfdoc.close()
    return {"pdf_sha256": digest(pdf), "docx_sha256": digest(docx),
            "pages": pages, "page_count": len(pages),
            "word_page_review_basis": "These pages render the DOCX in LibreOffice; "
                                      "other Word applications may paginate differently.",
            "complete": False}


def render(packet_path, output, soffice=None, allow_library_output=False,
           prepared_by=None, author=None):
    result = validate(packet_path)
    if not result["valid"]:
        raise PacketError("\n".join(result["errors"]))
    optional("docx")
    optional("pypdfium2")
    prepared_by, author = resolve_branding(prepared_by, author)
    binary = find_soffice(soffice)
    if not binary or not Path(binary).is_file() or not os.access(binary, os.X_OK):
        raise Blocked("Installed Word renderer unavailable; provide --soffice or set " + SOFFICE_ENV)
    binary_path = Path(binary).resolve()
    if (Path(binary).name.lower() not in SOFFICE_NAMES or
            binary_path.is_relative_to(result["library"]) or
            binary_path.is_relative_to(result["assets"]) or
            binary_path.is_relative_to(result["manifest_path"].parent)):
        raise PacketError("Renderer must be an installed office executable outside source/input roots")
    output = Path(output)
    # Never follow an output symlink or overwrite an existing output tree.
    if output.exists() or output.is_symlink():
        raise PacketError("Output already exists; choose a new output directory")
    for parent in output.absolute().parents:
        if parent.is_symlink():
            raise PacketError("Symlink output ancestor refused")
    output = output.resolve()
    if output.is_relative_to(result["library"]):
        if not allow_library_output:
            raise PacketError("Output inside source library needs --allow-library-output")
        exclusions = [safe_path(result["library"], x, False).resolve()
                      for x in result["packet"].get("source_exclusions", [])]
        if not any(output.is_relative_to(e) for e in exclusions):
            raise PacketError("Authorized library output must be covered by source_exclusions")
        for src in result["packet"]["sources"]:
            if safe_path(result["library"], src["path"]).is_relative_to(output):
                raise PacketError("Authorized output subdirectory must be excluded from sources")
    output.mkdir(parents=True, mode=0o700)
    internal = output / "_internal"
    internal.mkdir(mode=0o700)
    json_write(internal / "source-claim-ledger.json", result["packet"])
    states = []
    for doc in result["documents"]:
        spec, did = doc["spec"], doc["spec"]["id"]
        folder = output / did
        folder.mkdir()
        refs, projection = citation_projection(result, doc)
        canonical_body = doc["markdown"]
        for kind, data in doc["blocks"]:
            if kind == "image":
                name = data[1].name
                asset_folder = folder / "assets"
                asset_folder.mkdir(exist_ok=True)
                copied = asset_folder / name
                if copied.exists() and digest(copied) != digest(data[1]):
                    raise PacketError("Two diagram assets have conflicting basenames")
                if not copied.exists():
                    shutil.copyfile(data[1], copied)
                canonical_body = re.sub(r"!\[" + re.escape(data[0]) + r"\]\([^)]+\)",
                                        lambda _: f"![{data[0]}](assets/{name})", canonical_body)
        canonical = canonical_body.rstrip() + "\n\n" + refs.replace("# References", "## References", 1)
        (folder / f"{did}.md").write_text(canonical, encoding="utf-8")
        (folder / "citations.md").write_text(refs, encoding="utf-8")
        json_write(folder / "citations.json", projection)
        word = folder / f"{did}.docx"
        try:
            write_docx(doc, canonical, result["logo"], word, folder, prepared_by, author)
        except Exception as e:
            raise Blocked(f"DOCX construction failed: {e}") from e
        with tempfile.TemporaryDirectory(prefix="word-render-", dir=internal) as td:
            proc = subprocess.run([str(binary), "-env:UserInstallation=" + Path(td).as_uri(),
                                   "--headless", "--convert-to", "pdf", "--outdir",
                                   str(folder), str(word)], capture_output=True, text=True,
                                  timeout=120, check=False)
        pdf = folder / f"{did}.pdf"
        if proc.returncode or not pdf.exists() or pdf.stat().st_size == 0:
            raise Blocked("Word-to-PDF export failed: " + proc.stderr[:1000] + proc.stdout[:1000])
        review_folder = internal / "reviews" / did
        review_folder.mkdir(parents=True)
        try:
            inspection = inspect_pdf(pdf, word, review_folder, prepared_by)
        except Exception as e:
            raise Blocked(f"Rendered-page inspection failed: {e}") from e
        inspection.update({"id": did, "audience": spec["audience"],
                           "markdown_sha256": digest(folder / f"{did}.md"),
                           "citations_sha256": digest(folder / "citations.json"),
                           "citations_md_sha256": digest(folder / "citations.md"),
                           "canonical_content_sha256": hashlib.sha256(canonical.encode()).hexdigest(),
                           "estimated_grade": doc["estimated_grade"],
                           "logo_sha256": digest(result["logo"])})
        json_write(review_folder / "review.json", inspection)
        states.append(inspection)
    summary = {"schema_version": 1, "complete": False, "documents": states,
               "manifest_sha256": digest(packet_path),
               "review_required": "Every final page needs explicit visual attestation."}
    json_write(output / "packet-state.json", summary)
    return summary


def review_complete(output, attestation_path):
    output = Path(output).resolve()
    state = json_read(output / "packet-state.json")
    attestation = json_read(attestation_path)
    if not all(attestation.get(k) for k in ("reviewer", "inspected_at")):
        raise PacketError("Named reviewer and inspection time required")
    supplied = {d["id"]: d for d in attestation.get("documents", [])}
    if len(supplied) != len(attestation.get("documents", [])):
        raise PacketError("Duplicate document attestation")
    if set(supplied) != {d["id"] for d in state["documents"]}:
        raise PacketError("Attestation must cover every document, with no extras")
    for doc in state["documents"]:
        did = doc["id"]
        folder = safe_path(output, did, False)
        review_folder = safe_path(output, "_internal/reviews/" + did, False)
        a = supplied[did]
        for ext, key in (("pdf", "pdf_sha256"), ("docx", "docx_sha256"),
                         ("md", "markdown_sha256")):
            actual = digest(safe_path(folder, f"{did}.{ext}"))
            if actual != doc[key] or a.get(key) != actual:
                raise PacketError(f"{did}: stale {ext} review hash")
        if digest(safe_path(folder, "citations.json")) != doc["citations_sha256"]:
            raise PacketError(f"{did}: citation projection changed after render")
        if digest(safe_path(folder, "citations.md")) != doc["citations_md_sha256"]:
            raise PacketError(f"{did}: Markdown citation list changed after render")
        pages = {x["page"]: x for x in a.get("pages", [])}
        if len(pages) != len(a.get("pages", [])) or set(pages) != set(range(1, doc["page_count"] + 1)):
            raise PacketError(f"{did}: every page requires exactly one attestation")
        for page in doc["pages"]:
            record = pages[page["page"]]
            png_hash = digest(safe_path(review_folder, page["png"]))
            if png_hash != page["png_sha256"] or record.get("png_sha256") != png_hash:
                raise PacketError(f"{did}: stale page image review")
            if record.get("visual_pass") is not True or not all(
                    record.get("checks", {}).get(k) is True for k in PAGE_CHECKS):
                raise PacketError(f"{did}: page visual checks incomplete")
            if (page["out_of_page_glyphs"] or page.get("out_of_page_images") or page["missing_glyphs"] or
                    not page["footer_text_present"] or not page["footer_logo_detected"] or
                    page["footer_region_extra_text"]):
                raise PacketError(f"{did}: unresolved geometry/footer/glyph finding")
    final = {"schema_version": 1, "complete": True, "reviewer": attestation["reviewer"],
             "inspected_at": attestation["inspected_at"],
             "state_sha256": digest(output / "packet-state.json"),
             "attestation_sha256": digest(attestation_path), "documents": state["documents"],
             "limit": "Supplied attestations record review, not an automatic semantic or visual guarantee."}
    json_write(output / "review-complete.json", final)
    return final


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="command", required=True)
    inv = sub.add_parser("inventory")
    inv.add_argument("library")
    inv.add_argument("--exclude", action="append", default=[])
    inv.add_argument("--output", required=True)
    val = sub.add_parser("validate")
    val.add_argument("packet")
    ren = sub.add_parser("render")
    ren.add_argument("packet")
    ren.add_argument("--output", required=True)
    ren.add_argument("--soffice", help="LibreOffice executable (default: $SOFFICE_PATH, PATH, then Windows install paths)")
    ren.add_argument("--prepared-by", help="Footer preparer text (default: $LIBRARY_PACKETS_PREPARED_BY or Legion Code Inc.)")
    ren.add_argument("--author", help="DOCX author property (default: $LIBRARY_PACKETS_AUTHOR or the prepared-by text)")
    ren.add_argument("--allow-library-output", action="store_true")
    rev = sub.add_parser("review-complete")
    rev.add_argument("output")
    rev.add_argument("--attestation", required=True)
    args = ap.parse_args(argv)
    try:
        if args.command == "inventory":
            data = inventory(args.library, args.exclude)
            json_write(args.output, data)
        elif args.command == "validate":
            r = validate(args.packet)
            data = {k: r[k] for k in ("valid", "errors", "warnings")}
            print(json.dumps(data, indent=2))
            return 0 if r["valid"] else 1
        elif args.command == "render":
            data = render(args.packet, args.output, args.soffice, args.allow_library_output,
                          args.prepared_by, args.author)
        else:
            data = review_complete(args.output, args.attestation)
        print(json.dumps(data, indent=2))
        return 0
    except (PacketError, OSError, ValueError, KeyError, TypeError, AttributeError, subprocess.TimeoutExpired) as e:
        print(json.dumps({"complete": False, "status": "blocked" if isinstance(e, Blocked) else "invalid",
                          "error": str(e)}), file=sys.stderr)
        return 2 if isinstance(e, Blocked) else 1


if __name__ == "__main__":
    sys.exit(main())
