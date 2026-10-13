"""Behavioral checks for generic investigation helpers and generated Word content."""
import sys
sys.dont_write_bytecode = True
import importlib.util,json,tempfile,subprocess,os,zipfile
from pathlib import Path
from xml.etree import ElementTree
HERE=Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

def module(name):
 spec=importlib.util.spec_from_file_location(name,HERE/(name+'.py'))
 out=importlib.util.module_from_spec(spec);spec.loader.exec_module(out);return out

def main():
 count=0
 def check(label,fn):
  nonlocal count
  fn();count+=1;print('PASS',label)
 with tempfile.TemporaryDirectory(prefix='forensic-helpers-') as tmp:
  root=Path(tmp);out=root/'extracts';pdf=root/'pdf';pdf.mkdir()
  messages={'have_date':[{'_from_email':'billing@vendor.example','subject':'Customer Invoice #1001','_body':'Invoice #1001\nAmount Due: $120.00\nInvoice Items\nmonthly support $120.00 USD\nSubtotal','_dt':'2026-01-10','_id':'M-01'},{'_from_email':'other@other.example','subject':'Customer Invoice #9999','_body':'Invoice #9999\nAmount Due: $888.00','_dt':'2026-01-10','_id':'M-99'}]}
  inp=root/'processed.json';inp.write_text(json.dumps(messages))
  def run(name,args):
   p=subprocess.run([sys.executable,'-B',str(HERE/name),*map(str,args)],capture_output=True,text=True)
   assert p.returncode==0,p.stderr;return p
  check('explicit vendor and sender filters',lambda:run('parse_invoices.py',['--processed-json',inp,'--pdf-dir',pdf,'--email-vendor','Fictional email vendor','--email-senders','billing@vendor.example','--pdf-vendor','Fictional PDF vendor','--out',out]))
  invoices=json.loads((out/'email_invoices.json').read_text())
  assert len(invoices)==1 and invoices[0]['vendor']=='Fictional email vendor'
  check('missing vendor refuses before output',lambda:missing_vendor(inp,pdf,root))
  # Review-mode monthly model: endpoints with a single missing month.
  data=[{'invoice_number':'INV-1','issued':'2026-01-10','date_received':'2026-01-10','vendor':'Fictional email vendor','amount':120,'items':[{'description':'monthly support','amount':120}]},{'invoice_number':'INV-2','issued':'2026-03-10','date_received':'2026-03-10','vendor':'Fictional email vendor','amount':120,'items':[{'description':'monthly support','amount':120}]}]
  email=root/'email.json';email.write_text(json.dumps(data));pdfinput=root/'pdf.json';pdfinput.write_text('[]');derived=root/'derived.json'
  check('generic recurring gap is visibly derived',lambda:run('extrapolate_recurring.py',['--email-invoices',email,'--pdf-invoices',pdfinput,'--out',derived]))
  estimates=json.loads(derived.read_text());assert len(estimates)==1 and estimates[0]['date']=='2026-02-10' and estimates[0]['is_extrapolated']
  changed=json.loads(email.read_text());changed[1]['items'][0]['amount']=150
  changefile=root/'price-change.json';changefile.write_text(json.dumps(changed));changeout=root/'price-change-output.json'
  def refuse_prices():
   p=subprocess.run([sys.executable,'-B',str(HERE/'extrapolate_recurring.py'),'--email-invoices',str(changefile),'--pdf-invoices',str(pdfinput),'--out',str(changeout)],capture_output=True,text=True)
   assert p.returncode!=0 and 'multiple observed prices' in p.stderr and not changeout.exists()
  check('ambiguous price eras refuse without output',refuse_prices)
  wb=root/'book'/'invoices.xlsx'
  check('generic variable-size workbook generates',lambda:run('build_invoice_xlsx.py',['--email-invoices',email,'--pdf-invoices',pdfinput,'--extrapolated',derived,'--project-name','Fictional review','--out',wb]))
  from openpyxl import load_workbook
  book=load_workbook(wb,data_only=False);assert len(book.sheetnames)==3
  values=[str(c.value) for sheet in book for row in sheet for c in row if c.value is not None]
  assert any('Fictional email vendor documented' in v for v in values)
  # Make sure untrusted invoice text remains inert in XLSX.
  data[0]['items'][0]['description']='=HYPERLINK("https://example.invalid","bad")';email.write_text(json.dumps(data));wb2=root/'book'/'inert.xlsx'
  check('formula-like invoice text remains inert',lambda:run('build_invoice_xlsx.py',['--email-invoices',email,'--pdf-invoices',pdfinput,'--extrapolated',derived,'--project-name','Fictional review','--out',wb2]))
  book2=load_workbook(wb2,data_only=False);assert not any(c.data_type=='f' for sh in book2 for row in sh for c in row)
  git=module('parse_git_log');check('Git effort remains an explicit heuristic',lambda:assert_estimate(git))
  check('Word package parts are well-formed XML',lambda:check_word_documents())
 print(json.dumps({'passed':count}))

def missing_vendor(inp,pdf,root):
 p=subprocess.run([sys.executable,'-B',str(HERE/'parse_invoices.py'),'--processed-json',str(inp),'--pdf-dir',str(pdf),'--out',str(root/'missing')],capture_output=True,text=True)
 assert p.returncode!=0 and not (root/'missing').exists()

def assert_estimate(git):
 value=git.estimate_hours({'subject':'feature','files_changed':3,'insertions':30,'deletions':0})
 assert value==1

def check_word_documents():
 node=os.environ.get('FORENSIC_NODE','node')
 with tempfile.TemporaryDirectory(prefix='forensic-word-') as tmp:
  for entry in ['master_report','agency_report','attorney_memo','plain_language','pre_litigation']:
   out=Path(tmp)/entry
   p=subprocess.run([node,str(HERE/f'build_{entry}.js'),'--case-facts',str((HERE/'../examples/example-case-a/case-facts.json').resolve()),'--out-dir',str(out)],capture_output=True,text=True)
   assert p.returncode==0,p.stderr
   f=Path(json.loads(p.stdout)['output'])
   with zipfile.ZipFile(f) as z:
    assert z.testzip() is None
    for name in z.namelist():
     if name.endswith('.xml'):ElementTree.fromstring(z.read(name))
    body=z.read('word/document.xml').decode()
    assert 'FICTIONAL TRAINING EXAMPLE' in body and 'ESTIMATE' in body and 'HYPOTHESIS' in body and 'E-01' in body

if __name__=='__main__':main()
